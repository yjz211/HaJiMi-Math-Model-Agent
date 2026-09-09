"""Render HaJiMi's generated AI-report DOCX template using managed PyMuPDF.

This is a template-specific reflow exporter, not a general Word layout engine.
All paragraph/table text is consumed from the saved DOCX and checked in the PDF.
"""
from __future__ import annotations

import hashlib
import html
import json
from pathlib import Path
import re
import tempfile
from collections import Counter

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
import pymupdf as fitz


def normalized(text):
    return re.sub(r"\s+", "", text).replace("\u00ad", "")


def export_generated_report(docx_path: Path, pdf_path: Path) -> dict:
    if pdf_path.exists():
        raise FileExistsError("Refusing to overwrite an existing AI report PDF")
    document = Document(docx_path)
    if document.core_properties.title != "AI工具使用详情":
        raise ValueError("Only the generated HaJiMi AI report template is supported")
    if document.element.body.xpath(".//w:drawing | .//w:object | .//w:ins | .//w:del | .//w:hyperlink"):
        raise ValueError("Unsupported DOCX content; do not silently drop drawings, revisions or hyperlinks")
    pieces, texts = [], []

    def paragraph(p):
        texts.append(p.text)
        runs = []
        for run in p.runs:
            text = html.escape(run.text).replace("\n", "<br/>")
            if run.bold:
                text = f"<b>{text}</b>"
            if run.italic:
                text = f"<i>{text}</i>"
            runs.append(text)
        content = "".join(runs)
        name = p.style.name
        tag = "h1" if name == "Title" else "h2" if name == "Heading 1" else "h3" if name == "Heading 2" else "p"
        style = ' style="page-break-before:always"' if p.paragraph_format.page_break_before else ""
        return f"<{tag}{style}>{content}</{tag}>"

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            pieces.append(paragraph(block))
        elif isinstance(block, Table):
            widths = [cell.width or 1 for cell in block.rows[0].cells]
            total = sum(widths)
            rows = []
            for row in block.rows:
                header = bool(row._tr.xpath("./w:trPr/w:tblHeader"))
                tag = "th" if header else "td"
                cells = []
                for index, cell in enumerate(row.cells):
                    if len(cell.tables):
                        raise ValueError("Nested tables are unsupported in the AI report template")
                    content = "".join(paragraph(p) for p in cell.paragraphs)
                    cells.append(f'<{tag} style="width:{widths[index] / total * 100:.2f}%">{content}</{tag}>')
                rows.append("<tr>" + "".join(cells) + "</tr>")
            pieces.append("<table>" + "".join(rows) + "</table>")
    full_text = "\n".join(texts)
    if "工作草稿" in full_text or re.search(r"【待补充[^】]*】|\b(?:TODO|TBD)\b", full_text, re.I):
        raise ValueError("Draft or incomplete AI reports cannot be exported as submission PDF")
    css = """
      body { font-family: sans-serif; font-size: 11pt; line-height: 1.45; color: #000; }
      h1 { font-size: 19pt; text-align:center; margin: 0 0 14pt; }
      h2 { font-size:14pt; margin:14pt 0 7pt; page-break-after:avoid; }
      h3 { font-size:12pt; margin:10pt 0 5pt; page-break-after:avoid; }
      p { margin: 0 0 6pt; }
      table { border-collapse:collapse; width:100%; margin:7pt 0 10pt; font-size:10pt; }
      th,td { border:0.5pt solid #d9d9d9; padding:5pt; vertical-align:middle; }
      th { font-weight:bold; }
      td p,th p { margin:0; }
    """
    story = fitz.Story(html="<html><body>" + "".join(pieces) + "</body></html>", user_css=css)
    with tempfile.TemporaryDirectory(prefix="ai-report-render-", dir=pdf_path.parent) as raw:
        temporary = Path(raw) / "render.pdf"
        writer = fitz.DocumentWriter(str(temporary))
        try:
            for _ in range(100):
                device = writer.begin_page(fitz.Rect(0, 0, 595.276, 841.89))
                more, filled = story.place(fitz.Rect(45, 45, 550.276, 790))
                story.draw(device)
                writer.end_page()
                if not more:
                    break
            else:
                raise ValueError("Report exceeds 100 pages or contains a table that cannot fit")
        finally:
            writer.close()
            del writer
        with fitz.open(temporary) as pdf:
            rendered = normalized("".join(page.get_text(sort=True) for page in pdf))
            # PDF extractors interleave adjacent table columns at equal y positions.
            # Compare every non-whitespace character, not invalid cell substrings.
            expected_chars = Counter(normalized(full_text))
            actual_chars = Counter(rendered)
            if expected_chars != actual_chars:
                raise ValueError(f"DOCX/PDF character mismatch: missing={expected_chars - actual_chars}, extra={actual_chars - expected_chars}")
            for index, page in enumerate(pdf):
                if not page.get_text().strip():
                    raise ValueError("AI report contains a blank page")
                for block in page.get_text("dict")["blocks"]:
                    for line in block.get("lines", []):
                        for span in line["spans"]:
                            if not fitz.Rect(span["bbox"]) in page.rect:
                                raise ValueError("AI report text extends outside its page")
                page.insert_text((285, 817), str(index + 1), fontsize=9)
            pdf.set_metadata({"title": "AI工具使用详情"})
            pdf.save(pdf_path, garbage=4, deflate=True)
            count = len(pdf)
    receipt = {"docx_sha256": hashlib.sha256(docx_path.read_bytes()).hexdigest(),
               "pdf_sha256": hashlib.sha256(pdf_path.read_bytes()).hexdigest(), "pages": count,
               "renderer": "managed-pymupdf-generated-docx-template-reflow", "text_checks": len(texts),
               "visual_review": "pending"}
    (pdf_path.parent / "ai-report-receipt.json").write_text(json.dumps(receipt, indent=2), encoding="utf-8")
    return receipt
