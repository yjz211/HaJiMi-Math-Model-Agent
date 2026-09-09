#!/usr/bin/env python3
"""Build an honest, anonymous AI usage report as editable Word and optional PDF."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor
from lxml import etree


PLACEHOLDER = "【待补充】"
FINAL_STATUSES = {"final", "提交版"}
ADOPTION_VALUES = {"未采纳", "部分采纳", "采纳"}
GENERIC_VERIFICATIONS = {"已人工检查", "已核对", "结果正确", "人工复核", "已验证"}
ACTION_WORDS = (
    "对照", "查阅", "推导", "量纲", "逐行", "测试", "计算", "复算", "运行",
    "误差", "扰动", "一致性", "样例", "复现", "核查", "比较",
)


def value(data: dict, key: str) -> str:
    text = str(data.get(key, "")).strip()
    return text or PLACEHOLDER


def contains_placeholder(item: object) -> bool:
    if isinstance(item, dict):
        return any(contains_placeholder(v) for v in item.values())
    if isinstance(item, list):
        return any(contains_placeholder(v) for v in item)
    text = str(item or "")
    return bool(re.search(r"【待补充[^】]*】", text)) or bool(re.search(r"\b(?:TODO|TBD)\b", text, re.I))


def validate(data: dict) -> tuple[bool, list[str]]:
    status = str(data.get("document_status", "draft")).strip().lower()
    final = status in FINAL_STATUSES
    errors: list[str] = []
    tools = data.get("tools")
    records = data.get("records")
    if not isinstance(tools, list) or not tools:
        errors.append("tools 必须是非空列表")
        tools = []
    if not isinstance(records, list) or not records:
        errors.append("records 必须是非空列表")
        records = []

    tool_keys = ("name", "version", "provider", "use_method", "date_range")
    record_keys = (
        "id", "tool", "stage", "purpose", "provided_materials", "prompt_summary",
        "process", "ai_output", "paper_location", "adoption", "human_modification",
        "verification",
    )
    for index, tool in enumerate(tools, 1):
        if not isinstance(tool, dict):
            errors.append(f"第 {index} 个工具不是对象")
            continue
        for key in tool_keys:
            if not str(tool.get(key, "")).strip():
                errors.append(f"第 {index} 个工具缺少 {key}")
            elif not final and contains_placeholder(tool.get(key)):
                errors.append(f"第 {index} 个工具需补充 {key}")

    expected_ids = [str(index) for index in range(1, len(records) + 1)]
    actual_ids: list[str] = []
    tool_names = {str(tool.get("name", "")).strip() for tool in tools if isinstance(tool, dict)}
    for index, record in enumerate(records, 1):
        if not isinstance(record, dict):
            errors.append(f"第 {index} 条记录不是对象")
            continue
        for key in record_keys:
            if not str(record.get(key, "")).strip():
                errors.append(f"第 {index} 条记录缺少 {key}")
            elif not final and contains_placeholder(record.get(key)):
                errors.append(f"第 {index} 条记录需补充 {key}")
        actual_ids.append(str(record.get("id", "")).strip())
        if final and record.get("tool") not in tool_names:
            errors.append(f"第 {index} 条记录的工具未出现在 tools 中")
        if final and record.get("adoption") not in ADOPTION_VALUES:
            errors.append(f"第 {index} 条记录的 adoption 必须为未采纳、部分采纳或采纳")
        verification = str(record.get("verification", "")).strip().rstrip("。；; ")
        if final and (verification in GENERIC_VERIFICATIONS or not any(word in verification for word in ACTION_WORDS)):
            errors.append(f"第 {index} 条记录缺少可复核的核验动作")
    if actual_ids and actual_ids != expected_ids:
        errors.append(f"记录编号必须连续为 {expected_ids}")

    if final and contains_placeholder(data):
        errors.append("最终提交版不能包含待补充或 TODO/TBD 占位内容")
    if final and not str(data.get("responsibility_statement", "")).strip():
        errors.append("最终提交版缺少 responsibility_statement")
    return final, errors


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, value_twips: int = 90) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge in ("top", "left", "bottom", "right"):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value_twips))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9", size: str = "6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_run_font(run, east_asia: str = "宋体", western: str = "Arial", size: float = 10.5) -> None:
    run.font.name = western
    run.font.size = Pt(size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instr, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=9)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths_cm: list[float]):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_header(header_row)
    set_cant_split(header_row)
    for index, (cell, label) in enumerate(zip(header_row.cells, headers)):
        cell.width = Cm(widths_cm[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
        set_cell_shading(cell, "2E5F8A")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(label)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_run_font(run, east_asia="黑体", size=9.5)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        set_cant_split(row)
        cells = row.cells
        for index, (cell, text) in enumerate(zip(cells, values)):
            cell.width = Cm(widths_cm[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, "F4F7FA")
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(str(text))
            set_run_font(run, size=9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_detail_table(doc: Document, fields: list[tuple[str, str]]):
    """Render one AI-use record as a readable vertical field table."""
    table = doc.add_table(rows=0, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    for label, content in fields:
        row = table.add_row()
        set_cant_split(row)
        label_cell, content_cell = row.cells
        label_cell.width = Cm(3.4)
        content_cell.width = Cm(13.2)
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, 110)
        set_cell_shading(label_cell, "2E5F8A")
        label_paragraph = label_cell.paragraphs[0]
        label_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        label_paragraph.paragraph_format.space_after = Pt(0)
        label_run = label_paragraph.add_run(label)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor(255, 255, 255)
        set_run_font(label_run, east_asia="黑体", size=10)
        content_paragraph = content_cell.paragraphs[0]
        content_paragraph.paragraph_format.space_after = Pt(0)
        content_paragraph.paragraph_format.line_spacing = 1.2
        content_run = content_paragraph.add_run(content)
        set_run_font(content_run, size=10.5)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)
    return table


def add_labeled_paragraph(doc: Document, label: str, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.line_spacing = 1.25
    label_run = paragraph.add_run(f"{label}：")
    label_run.bold = True
    set_run_font(label_run, east_asia="黑体", size=10.5)
    content_run = paragraph.add_run(text)
    set_run_font(content_run, size=10.5)


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(22)
    section.bottom_margin = Mm(20)
    section.left_margin = Mm(22)
    section.right_margin = Mm(22)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.space_after = Pt(5)

    title = doc.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size in (("Heading 1", 14), ("Heading 2", 12)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])


def scrub_metadata(docx_path: Path) -> None:
    """Remove personal core properties and Word revision-session identifiers."""
    word_namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    rel_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    content_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    with tempfile.NamedTemporaryFile(dir=docx_path.parent, suffix=".docx", delete=False) as stream:
        temp_path = Path(stream.name)
    try:
        with zipfile.ZipFile(docx_path, "r") as source, zipfile.ZipFile(temp_path, "w", zipfile.ZIP_DEFLATED) as target:
            for info in source.infolist():
                if info.filename == "docProps/custom.xml":
                    continue
                content = source.read(info.filename)
                if info.filename in {"[Content_Types].xml", "_rels/.rels"}:
                    try:
                        root = etree.fromstring(content)
                        for element in list(root):
                            is_custom_type = (
                                element.tag == f"{{{content_namespace}}}Override"
                                and element.get("PartName", "") == "/docProps/custom.xml"
                            )
                            is_custom_relation = (
                                element.tag == f"{{{rel_namespace}}}Relationship"
                                and element.get("Target", "").endswith("docProps/custom.xml")
                            )
                            if is_custom_type or is_custom_relation:
                                root.remove(element)
                        content = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone="yes")
                    except etree.XMLSyntaxError:
                        pass
                elif info.filename == "docProps/core.xml" or (
                    info.filename.startswith("word/") and info.filename.endswith(".xml")
                ):
                    try:
                        root = etree.fromstring(content)
                        if info.filename == "docProps/core.xml":
                            for xpath, namespaces in (
                                (".//dc:creator", {"dc": "http://purl.org/dc/elements/1.1/"}),
                                (".//cp:lastModifiedBy", {"cp": "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"}),
                            ):
                                for element in root.xpath(xpath, namespaces=namespaces):
                                    element.text = ""
                        else:
                            for element in list(root.iter()):
                                for attribute in list(element.attrib):
                                    if attribute.startswith(f"{{{word_namespace}}}rsid"):
                                        del element.attrib[attribute]
                                if element.tag in {
                                    f"{{{word_namespace}}}rsids",
                                    f"{{{word_namespace}}}rsidRoot",
                                    f"{{{word_namespace}}}rsid",
                                }:
                                    parent = element.getparent()
                                    if parent is not None:
                                        parent.remove(element)
                        content = etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone="yes")
                    except etree.XMLSyntaxError:
                        pass
                target.writestr(info.filename, content)
        temp_path.replace(docx_path)
    finally:
        temp_path.unlink(missing_ok=True)


def build_report(data: dict, output_docx: Path, final: bool) -> None:
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.author = ""
    props.last_modified_by = ""
    props.title = "AI工具使用详情"
    props.subject = ""
    props.keywords = ""
    props.comments = ""

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI工具使用详情")
    if not final:
        draft = doc.add_paragraph()
        draft.alignment = WD_ALIGN_PARAGRAPH.CENTER
        draft_run = draft.add_run("工作草稿  不可直接提交")
        draft_run.bold = True
        draft_run.font.color.rgb = RGBColor(180, 0, 0)
        set_run_font(draft_run, east_asia="黑体", size=11)

    tools = data.get("tools", [])
    records = data.get("records", [])

    doc.add_heading("一、所用AI工具名称、版本或型号", level=1)
    add_table(
        doc,
        ["序号", "AI工具", "版本或型号", "提供方", "使用方式", "使用日期"],
        [[
            str(index), value(tool, "name"), value(tool, "version"), value(tool, "provider"),
            value(tool, "use_method"), value(tool, "date_range"),
        ] for index, tool in enumerate(tools, 1)],
        [1.0, 2.7, 3.8, 2.3, 2.8, 3.0],
    )

    doc.add_heading("二、具体使用目的和环节", level=1)
    add_table(
        doc,
        ["记录编号", "AI工具", "使用环节", "具体使用目的", "对应位置"],
        [[value(record, "id"), value(record, "tool"), value(record, "stage"),
          value(record, "purpose"), value(record, "paper_location")] for record in records],
        [1.7, 2.7, 2.8, 4.8, 4.6],
    )

    doc.add_heading("三、主要提示方式与使用过程说明", level=1)
    for index, record in enumerate(records, 1):
        doc.add_heading(f"3.{index} 使用记录 {value(record, 'id')}：{value(record, 'stage')}", level=2)
        for label, key in (
            ("使用工具", "tool"), ("使用目的", "purpose"),
            ("提供给AI的材料", "provided_materials"), ("主要提示方式", "prompt_summary"),
            ("使用过程", "process"), ("AI输出概述", "ai_output"),
            ("对应位置", "paper_location"),
        ):
            add_labeled_paragraph(doc, label, value(record, key))

    verification_heading = doc.add_heading("四、AI输出的采纳、人工修改和核验情况", level=1)
    verification_heading.paragraph_format.page_break_before = True
    for index, record in enumerate(records, 1):
        doc.add_heading(f"4.{index} 记录 {value(record, 'id')}", level=2)
        add_detail_table(doc, [
            ("AI主要输出", value(record, "ai_output")),
            ("采纳情况", value(record, "adoption")),
            ("人工修改", value(record, "human_modification")),
            ("核验方式与结果", value(record, "verification")),
        ])

    examples = data.get("examples") or []
    if examples:
        doc.add_heading("五、典型交互示例", level=1)
        for index, example in enumerate(examples, 1):
            doc.add_heading(f"5.{index} 交互示例 {index}", level=2)
            for label, key in (
                ("对应记录", "record_id"), ("参赛队真实输入", "user_input"),
                ("AI关键输出或准确概括", "ai_output"), ("后续处理", "follow_up"),
            ):
                add_labeled_paragraph(doc, label, value(example, key))

    doc.add_heading("六、提交前检查与人工责任说明", level=1)
    doc.add_paragraph("以下为提交前待逐项确认的检查项，不表示本报告生成时已完成这些动作；实际核验人员、方法和结果以各条真实记录为准。")
    checks = [
        "工具、版本、用途和记录编号前后一致",
        "正文声明覆盖详情中的全部实质用途",
        "核验动作、结果及执行者有记录可查，AI或程序验证不写成人工核验",
        "未包含身份、账号、联系方式、内部路径或密钥",
        "无模板说明、空白示例和待补充占位符",
        "确认Word元数据清理结果和最终PDF已成功生成",
    ]
    for item in checks:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_run_font(run, size=10.5)
    add_labeled_paragraph(doc, "人工责任说明", value(data, "responsibility_statement"))

    output_docx.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_docx)
    scrub_metadata(output_docx)


def find_soffice(explicit: Path | None) -> Path:
    executable = "soffice.exe" if sys.platform == "win32" else "soffice"
    located = shutil.which(executable)
    candidates = [
        explicit,
        Path(located) if located else None,
        Path(r"C:\Program Files\LibreOffice\program\soffice.exe") if sys.platform == "win32" else None,
        Path(r"C:\Program Files (x86)\LibreOffice\program\soffice.exe") if sys.platform == "win32" else None,
    ]
    for candidate in candidates:
        if candidate and candidate.is_file():
            return candidate.resolve()
    raise FileNotFoundError("未找到 LibreOffice；请安装后重试，或用 --libreoffice 指定 soffice 路径")


def export_pdf(docx_path: Path, pdf_path: Path, soffice: Path) -> None:
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="ai_usage_word_") as temp_name:
        temp = Path(temp_name)
        out_dir = temp / "out"
        profile = temp / "profile"
        out_dir.mkdir()
        profile.mkdir()
        result = subprocess.run(
            [str(soffice), "--headless", f"-env:UserInstallation={profile.resolve().as_uri()}",
             "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )
        built = out_dir / f"{docx_path.stem}.pdf"
        if result.returncode or not built.is_file():
            detail = (result.stdout + "\n" + result.stderr).strip()
            raise RuntimeError(f"Word 导出 PDF 失败：{detail}")
        shutil.copy2(built, pdf_path)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("config", type=Path)
    parser.add_argument("--export-pdf", action="store_true", help="用受管PDF运行库从生成的Word模板导出PDF（内容保真重排）")
    parser.add_argument("--libreoffice", type=Path, help="soffice 可执行文件路径")
    args = parser.parse_args()

    config_path = args.config.resolve()
    data = json.loads(config_path.read_text(encoding="utf-8"))
    final, errors = validate(data)
    if errors:
        if final:
            raise ValueError("最终提交版校验失败：\n- " + "\n- ".join(errors))
        structural = [error for error in errors if "必须是非空列表" in error or "不是对象" in error]
        if structural:
            raise ValueError("配置结构错误：\n- " + "\n- ".join(structural))

    output_docx = Path(data.get("output_docx", "AI工具使用详情.docx"))
    if not output_docx.is_absolute():
        output_docx = (config_path.parent / output_docx).resolve()
    output_pdf = Path(data.get("output_pdf", "AI工具使用详情.pdf"))
    if not output_pdf.is_absolute():
        output_pdf = (config_path.parent / output_pdf).resolve()

    build_report(data, output_docx, final)
    if args.export_pdf:
        if not final:
            raise ValueError("工作草稿不可导出为可提交 PDF；补全信息并将 document_status 设为 final 后重试")
        sys.path.insert(0, str(Path(__file__).resolve().parent))
        from export_ai_report_pdf import export_generated_report
        export_generated_report(output_docx, output_pdf)

    result = {
        "status": "final" if final else "draft_not_for_submission",
        "docx": str(output_docx),
        "pdf": str(output_pdf) if args.export_pdf and final and output_pdf.is_file() else None,
        "validation_notes": errors,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
