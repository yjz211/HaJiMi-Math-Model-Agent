from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any


def inspect_input(path: Path) -> dict[str, Any]:
    path = path.resolve()
    if not path.is_file():
        raise ValueError(f"Input is not a regular file: {path}")
    suffix = path.suffix.lower()
    base: dict[str, Any] = {"path": path.name, "size_bytes": path.stat().st_size, "type": suffix.lstrip(".")}
    if suffix == ".csv":
        return {**base, **_inspect_csv(path)}
    if suffix == ".xlsx":
        return {**base, **_inspect_xlsx(path)}
    if suffix == ".pdf":
        return {**base, **_inspect_pdf(path)}
    if suffix == ".docx":
        return {**base, **_inspect_docx(path)}
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        return {**base, "json_kind": type(value).__name__, "top_level_size": len(value) if hasattr(value, "__len__") else None}
    if suffix in {".txt", ".md", ".tex"}:
        text = path.read_text(encoding="utf-8")
        return {**base, "line_count": len(text.splitlines()), "character_count": len(text)}
    return base


def _inspect_csv(path: Path) -> dict[str, Any]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        sample = handle.read(65536)
        handle.seek(0)
        dialect = csv.Sniffer().sniff(sample)
        reader = csv.reader(handle, dialect)
        rows = list(reader)
    return {
        "row_count": max(0, len(rows) - 1),
        "column_count": len(rows[0]) if rows else 0,
        "columns": rows[0] if rows else [],
        "delimiter": dialect.delimiter,
    }


def _inspect_xlsx(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        return {
            "sheets": [
                {"name": sheet.title, "max_row": sheet.max_row, "max_column": sheet.max_column}
                for sheet in workbook.worksheets
            ]
        }
    finally:
        workbook.close()


def _inspect_pdf(path: Path) -> dict[str, Any]:
    from pypdf import PdfReader

    reader = PdfReader(path)
    text_characters = sum(len(page.extract_text() or "") for page in reader.pages)
    return {"page_count": len(reader.pages), "extractable_text_characters": text_characters, "encrypted": reader.is_encrypted}


def _inspect_docx(path: Path) -> dict[str, Any]:
    from docx import Document

    document = Document(path)
    return {"paragraph_count": len(document.paragraphs), "table_count": len(document.tables)}

